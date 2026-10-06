from pathlib import Path
from datetime import datetime, timedelta

import numpy as np
import openpyxl
import rasterio
from rasterio.transform import from_origin


# ============================================================
# PATHS
# ============================================================

# Repository root:
# hydro-data-tools/
# ├─ src/
# ├─ input/
# └─ output/
REPO_ROOT = Path(__file__).resolve().parents[1]

INPUT_DIR = REPO_ROOT / "input"
OUTPUT_DIR = REPO_ROOT / "output"

HYETOGRAPH_FILE = INPUT_DIR / "Hyetograph.xlsx"
REFERENCE_RASTER = INPUT_DIR / "reference_extent.tif"

PRECIP_OUTPUT_DIR = OUTPUT_DIR / "precip_rasters"


# ============================================================
# SETTINGS
# ============================================================

RAINFALL_CELL_SIZE = 500.0  # metres
BUFFER = 500.0              # metres around reference raster extent


# ============================================================
# HMS DATETIME PARSING
# ============================================================

def parse_hms_datetime(value: str) -> datetime:
    """
    Parse HEC-HMS DSSVue-exported timestamps such as:

        01 Jan 2026, 00:10
        01 Jan 2026, 24:00

    Python datetime does not accept 24:00, so it is converted
    to 00:00 of the following day.
    """
    value = value.strip()

    date_part, time_part = [
        part.strip()
        for part in value.split(",", maxsplit=1)
    ]

    if time_part == "24:00":
        date = datetime.strptime(date_part, "%d %b %Y")
        return date + timedelta(days=1)

    return datetime.strptime(
        f"{date_part}, {time_part}",
        "%d %b %Y, %H:%M",
    )


# ============================================================
# READ HYETOGRAPH
# ============================================================

def read_hyetograph(file_path: Path):
    """
    Read the HMS Frequency Storm Excel export.

    Expected structure in the provided file:
        Row 2: headers
        Row 4: Units = MM
        Row 5: Type = PER-CUM
        Row 6 onward: precipitation records
    """
    if not file_path.exists():
        raise FileNotFoundError(
            f"Hyetograph file not found:\n{file_path}"
        )

    workbook = openpyxl.load_workbook(
        file_path,
        data_only=True,
        read_only=True,
    )

    worksheet = workbook.active

    units = worksheet.cell(row=4, column=3).value
    data_type = worksheet.cell(row=5, column=3).value

    units = str(units).strip().upper()
    data_type = str(data_type).strip().upper()

    if units != "MM":
        raise ValueError(
            f"Expected precipitation units 'MM', found '{units}'."
        )

    if data_type != "PER-CUM":
        raise ValueError(
            f"Expected DSS type 'PER-CUM', found '{data_type}'."
        )

    records = []

    for row in worksheet.iter_rows(
        min_row=6,
        values_only=True,
    ):
        _, date_time, precipitation = row[:3]

        if date_time is None:
            continue

        timestamp = parse_hms_datetime(str(date_time))

        precipitation_mm = (
            0.0 if precipitation is None
            else float(precipitation)
        )

        if precipitation_mm < 0:
            raise ValueError(
                f"Negative precipitation found at {timestamp}: "
                f"{precipitation_mm} mm"
            )

        records.append(
            {
                "end_time": timestamp,
                "precip_mm": precipitation_mm,
            }
        )

    workbook.close()

    if len(records) < 2:
        raise ValueError(
            "The hyetograph must contain at least two records."
        )

    return records


# ============================================================
# VALIDATE TIME SERIES
# ============================================================

def validate_hyetograph(records):
    timestep = (
        records[1]["end_time"]
        - records[0]["end_time"]
    )

    if timestep <= timedelta(0):
        raise ValueError("Invalid precipitation timestep.")

    for index in range(1, len(records)):
        current_timestep = (
            records[index]["end_time"]
            - records[index - 1]["end_time"]
        )

        if current_timestep != timestep:
            raise ValueError(
                "Inconsistent timestep detected between "
                f"{records[index - 1]['end_time']} and "
                f"{records[index]['end_time']}."
            )

    total_precipitation = sum(
        record["precip_mm"]
        for record in records
    )

    maximum_record = max(
        records,
        key=lambda record: record["precip_mm"],
    )

    return timestep, total_precipitation, maximum_record


# ============================================================
# CREATE RAINFALL GRID DEFINITION
# ============================================================

def create_grid_definition(reference_raster: Path):
    if not reference_raster.exists():
        raise FileNotFoundError(
            f"Reference raster not found:\n{reference_raster}"
        )

    with rasterio.open(reference_raster) as source:
        crs = source.crs
        bounds = source.bounds

    if crs is None:
        raise ValueError(
            "Reference raster has no CRS."
        )

    xmin = bounds.left - BUFFER
    ymin = bounds.bottom - BUFFER
    xmax = bounds.right + BUFFER
    ymax = bounds.top + BUFFER

    width = int(
        np.ceil(
            (xmax - xmin)
            / RAINFALL_CELL_SIZE
        )
    )

    height = int(
        np.ceil(
            (ymax - ymin)
            / RAINFALL_CELL_SIZE
        )
    )

    # Align actual maximum coordinates to the generated
    # grid dimensions.
    xmax = xmin + width * RAINFALL_CELL_SIZE
    ymax = ymin + height * RAINFALL_CELL_SIZE

    transform = from_origin(
        xmin,
        ymax,
        RAINFALL_CELL_SIZE,
        RAINFALL_CELL_SIZE,
    )

    return {
        "crs": crs,
        "transform": transform,
        "width": width,
        "height": height,
        "bounds": (xmin, ymin, xmax, ymax),
    }


# ============================================================
# CREATE GEOTIFFS
# ============================================================

def create_precipitation_rasters(
    records,
    timestep,
    grid,
):
    PRECIP_OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    for record in records:
        end_time = record["end_time"]
        start_time = end_time - timestep

        precipitation_mm = record["precip_mm"]

        data = np.full(
            (
                grid["height"],
                grid["width"],
            ),
            precipitation_mm,
            dtype=np.float32,
        )

        start_string = start_time.strftime(
            "%Y-%m-%dT%H%M"
        )

        end_string = end_time.strftime(
            "%Y-%m-%dT%H%M"
        )

        filename = (
            "PRECIPITATION_"
            f"{start_string}_to_{end_string}.tif"
        )

        output_path = (
            PRECIP_OUTPUT_DIR / filename
        )

        with rasterio.open(
            output_path,
            "w",
            driver="GTiff",
            height=grid["height"],
            width=grid["width"],
            count=1,
            dtype="float32",
            crs=grid["crs"],
            transform=grid["transform"],
            compress="deflate",
        ) as destination:
            destination.write(data, 1)


# ============================================================
# MAIN
# ============================================================

def main():
    print("Reading HEC-HMS hyetograph...")

    records = read_hyetograph(
        HYETOGRAPH_FILE
    )

    (
        timestep,
        total_precipitation,
        maximum_record,
    ) = validate_hyetograph(records)

    print()
    print("Hyetograph:")
    print(f"  Intervals:      {len(records)}")
    print(f"  Timestep:       {timestep}")
    print(
        f"  Total rainfall: "
        f"{total_precipitation:.3f} mm"
    )
    print(
        f"  Peak rainfall:  "
        f"{maximum_record['precip_mm']:.3f} mm"
    )
    print(
        f"  Peak end time:  "
        f"{maximum_record['end_time']}"
    )

    print()
    print("Reading reference raster...")

    grid = create_grid_definition(
        REFERENCE_RASTER
    )

    print(f"  CRS:        {grid['crs']}")
    print(
        f"  Resolution: "
        f"{RAINFALL_CELL_SIZE} m"
    )
    print(
        f"  Dimensions: "
        f"{grid['width']} x {grid['height']}"
    )

    print()
    print("Creating precipitation rasters...")

    create_precipitation_rasters(
        records,
        timestep,
        grid,
    )

    print()
    print(
        f"Created {len(records)} GeoTIFF files:"
    )
    print(f"  {PRECIP_OUTPUT_DIR}")


if __name__ == "__main__":
    main()